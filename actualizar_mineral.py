"""Lee el Excel INGRESO_MINERAL (hojas BD y BDSP) y genera mineral.json para la app.
Se ejecuta a mano o desde GitHub Actions (secreto DROPBOX_URL_MINERAL)."""
import io, json, os, sys, urllib.request, datetime
from openpyxl import load_workbook

URL = os.environ.get("DROPBOX_URL_MINERAL", "").strip()
SALIDA = "mineral.json"

def descargar(url):
    if not url: sys.exit("Falta el secreto DROPBOX_URL_MINERAL.")
    url = url.replace("dl=0", "dl=1")
    if "dropbox.com" in url and "dl=1" not in url:
        url += ("&" if "?" in url else "?") + "dl=1"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as r: return r.read()

def fecha(v):
    if isinstance(v, datetime.datetime): v = v.date()
    return v.isoformat() if isinstance(v, datetime.date) else None

def f(v):
    try: return float(v)
    except (TypeError, ValueError): return None

def leer(contenido):
    wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    for h in ("BD", "BDSP"):
        if h not in wb.sheetnames: sys.exit(f"No se encontró la hoja {h}.")
    # BDSP: lotes enviados al stock pile (columna C desde la fila 5)
    sp = set()
    for r in wb["BDSP"].iter_rows(min_row=5, min_col=3, max_col=3, values_only=True):
        try: sp.add(int(r[0]))
        except (TypeError, ValueError): pass
    # BD desde la fila 5: C lote, D fecha, E peso, F mina, K ley, L consumo, S peso chancado
    minas, idx, lotes = [], {}, []
    for r in wb["BD"].iter_rows(min_row=5, min_col=3, max_col=19, values_only=True):
        if r[0] is None: continue
        try: lote = int(r[0])
        except (TypeError, ValueError): lote = str(r[0]).strip()   # ej. "RECHAZO 16"
        fch, peso = fecha(r[1]), f(r[2])
        if fch is None or peso is None: continue
        mina = " ".join(str(r[3] or "SIN MINA").upper().split())
        if mina not in idx: idx[mina] = len(minas); minas.append(mina)
        ley, cons = f(r[8]), f(r[9])
        chanc = 1 if f(r[16]) else 0
        est = 1 if chanc else (2 if lote in sp else 0)   # 1 chancado, 2 stock pile, 0 disponible
        lotes.append([lote, fch, round(peso, 2), idx[mina], ley if ley else 0, cons if cons else 0, est])
    # BDCH: chancado por pila. C lote, D fecha chancado, E mina, G peso, H pila destino, J ley, K consumo, O precio
    pilas = []
    if "BDCH" in wb.sheetnames:
        for r in wb["BDCH"].iter_rows(min_row=5, min_col=3, max_col=15, values_only=True):
            if r[0] is None: continue
            try: lote = int(r[0])
            except (TypeError, ValueError): lote = str(r[0]).strip()
            fch, peso, pila = fecha(r[1]), f(r[4]), str(r[5] or "").strip().upper()
            if fch is None or peso is None or not pila: continue
            mina = " ".join(str(r[2] or "SIN MINA").upper().split())
            if mina not in idx: idx[mina] = len(minas); minas.append(mina)
            ley, cons = f(r[7]), f(r[8])
            precio = f(r[12])
            pilas.append([lote, fch, pila, idx[mina], round(peso, 2), ley if ley else 0, cons if cons else 0, round(precio) if precio else 0])
    # Tarifas ENAMI: hoja INFORME DE PILAS desde X14 (X fecha, AB base, AC escala, AD ácido)
    tarifas = []
    if "INFORME DE PILAS" in wb.sheetnames:
        for r in wb["INFORME DE PILAS"].iter_rows(min_row=14, min_col=24, max_col=30, values_only=True):
            fch = fecha(r[0])
            if fch is None or f(r[4]) is None: continue
            tarifas.append([fch[:7], f(r[4]), f(r[5]) or 0, f(r[6]) or 0])
    return minas, lotes, pilas, tarifas

def main(contenido=None):
    minas, lotes, pilas, tarifas = leer(contenido or descargar(URL))
    if len(lotes) < 50: sys.exit(f"Muy pocos lotes ({len(lotes)}): no se actualiza.")
    lotes.sort(key=lambda x: (x[1], str(x[0])))
    pilas.sort(key=lambda x: (x[1], x[2], str(x[0])))
    datos = {"actualizado": max(l[1] for l in lotes), "minas": minas, "lotes": lotes, "pilas": pilas, "tarifas": sorted(tarifas)}
    nuevo = json.dumps(datos, separators=(",", ":"), ensure_ascii=False)
    viejo = open(SALIDA, encoding="utf-8").read() if os.path.exists(SALIDA) else ""
    if nuevo == viejo: print("Sin cambios en mineral."); return
    open(SALIDA, "w", encoding="utf-8").write(nuevo)
    print(f"mineral.json: {len(lotes)} lotes, {len(pilas)} lotes chancados en pilas, {len(minas)} minas")

if __name__ == "__main__":
    main(open(sys.argv[1], "rb").read() if len(sys.argv) > 1 else None)
