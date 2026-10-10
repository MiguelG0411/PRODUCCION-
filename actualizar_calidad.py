"""Lee el Excel CONTROL_DE_CALIDAD (hojas BD y BDS) y genera pallets.json para la app.
1 lote = 1 pallet. Se ejecuta a mano o desde GitHub Actions (secreto DROPBOX_URL_CALIDAD)."""
import io, json, os, sys, urllib.request, datetime
from openpyxl import load_workbook

URL = os.environ.get("DROPBOX_URL_CALIDAD", "").strip()
SALIDA = "pallets.json"

def descargar(url):
    if not url:
        sys.exit("Falta el secreto DROPBOX_URL_CALIDAD.")
    url = url.replace("dl=0", "dl=1")
    if "dropbox.com" in url and "dl=1" not in url:
        url += ("&" if "?" in url else "?") + "dl=1"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=90) as r:
        return r.read()

def fecha(v):
    if isinstance(v, datetime.datetime): v = v.date()
    return v.isoformat() if isinstance(v, datetime.date) else None

def leer(contenido):
    wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    for h in ("BD", "BDS"):
        if h not in wb.sheetnames: sys.exit(f"No se encontró la hoja {h}.")
    # BD: fila 5 encabezados. A lote, B fecha, F %Cu, G %humedad, L categoría, T estado
    lotes = []
    for f in wb["BD"].iter_rows(min_row=6, max_col=20, values_only=True):
        lote, fch = f[0], fecha(f[1])
        if lote is None or fch is None: continue
        try: lote = int(lote)
        except (TypeError, ValueError): continue
        try: cat = int(f[11])
        except (TypeError, ValueError): cat = 0
        est = 1 if str(f[19] or "").strip().upper() == "ALMACENADO" else 0
        lotes.append([lote, fch, cat, est])
    # BDS: fila 3 encabezados. C lote, D fecha salida, E camión
    salidas = []
    for f in wb["BDS"].iter_rows(min_row=4, min_col=3, max_col=5, values_only=True):
        lote, fch, cam = f
        fch = fecha(fch)
        if fch is None: continue
        try: lote = int(lote)
        except (TypeError, ValueError): lote = 0  # SIN LOTE
        salidas.append([lote, fch, str(cam or "").strip()])
    return lotes, salidas

def main(contenido=None):
    lotes, salidas = leer(contenido or descargar(URL))
    if len(lotes) < 50:
        sys.exit(f"Muy pocos lotes ({len(lotes)}): no se actualiza.")
    lotes.sort(key=lambda x: (x[1], x[0])); salidas.sort(key=lambda x: x[1])
    datos = {"tonPallet": 1.25, "actualizado": max(l[1] for l in lotes), "lotes": lotes, "salidas": salidas}
    nuevo = json.dumps(datos, separators=(",", ":"), ensure_ascii=False)
    viejo = open(SALIDA, encoding="utf-8").read() if os.path.exists(SALIDA) else ""
    if nuevo == viejo:
        print("Sin cambios en pallets."); return
    open(SALIDA, "w", encoding="utf-8").write(nuevo)
    print(f"pallets.json: {len(lotes)} lotes, {len(salidas)} salidas")

if __name__ == "__main__":
    if len(sys.argv) > 1: main(open(sys.argv[1], "rb").read())
    else: main()
