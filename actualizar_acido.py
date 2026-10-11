"""Lee el Excel CONTROL_DE_ACIDO (hoja BDR) y genera acido.json con el consumo diario de ácido.
Reglas (definidas por la planta):
  - Toneladas de cada estanque (TK) = valor medido x 11,2. Excepción: TK3 desde el 27-02-2026 = valor medido x 20.
  - Consumo del día en cada TK = lectura del día anterior - lectura del día (solo si es positivo).
    Si la lectura sube, llegó un camión de ácido: no se considera consumo.
  - Consumo del día = suma de los TK.
  - Desde el 27-02-2026, si el mismo día baja el TK4 y sube el TK3 se considera traspaso (no es consumo).
  - Los días en que llega un camión se estima el consumo con el promedio de los 3 días anteriores y 3 posteriores.
Se ejecuta a mano (python actualizar_acido.py archivo.xlsm) o desde GitHub Actions (secreto DROPBOX_URL_ACIDO)."""
import io, json, os, sys, urllib.request, datetime
from openpyxl import load_workbook

URL = os.environ.get("DROPBOX_URL_ACIDO", "").strip()
SALIDA = "acido.json"
CAMBIO_TK3 = "2026-02-26"
MAX_LECTURA = 3.5   # una lectura mayor se interpreta como coma olvidada (ej. 70 -> 0,70)

def descargar(url):
    if not url: sys.exit("Falta el secreto DROPBOX_URL_ACIDO.")
    url = url.replace("dl=0", "dl=1")
    if "dropbox.com" in url and "dl=1" not in url: url += ("&" if "?" in url else "?") + "dl=1"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as r: return r.read()

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

def factor(tk, fecha):
    return 20.0 if (tk == 3 and fecha > CAMBIO_TK3) else 11.2

def leer(contenido):
    wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    if "BDR" not in wb.sheetnames: sys.exit("No se encontró la hoja BDR.")
    filas = []
    # BDR desde la fila 6 (más reciente arriba): C fecha, D mes, E-I TK1..TK5
    for i, r in enumerate(wb["BDR"].iter_rows(min_row=6, min_col=3, max_col=9, values_only=True), start=6):
        if r[0] is None and all(v is None for v in r[2:7]): continue
        f = r[0]
        if isinstance(f, datetime.datetime): f = f.date()
        f = f.isoformat() if isinstance(f, datetime.date) and f.year >= 2000 else None
        filas.append([i, f, [num(v) for v in r[2:7]]])
    filas.reverse()   # de la más antigua a la más reciente (orden real de las lecturas)
    avisos = {"fechas": [], "comas": []}
    # fechas inválidas: día siguiente a la lectura anterior
    for k, (i, f, t) in enumerate(filas):
        if f is None and k > 0 and filas[k-1][1]:
            f = (datetime.date.fromisoformat(filas[k-1][1]) + datetime.timedelta(days=1)).isoformat()
            filas[k][1] = f; avisos["fechas"].append(i)
        for j in range(5):
            if t[j] is not None and t[j] > MAX_LECTURA:
                avisos["comas"].append([i, j+1, t[j]]); t[j] = round(t[j] / 100, 3)
    dias, prev = [], [None]*5
    for i, f, t in filas:
        if f is None: continue
        cons, rec, nivel, sube3 = [0.0]*5, 0.0, [], 0.0
        for j in range(5):
            fac = factor(j+1, f)
            if t[j] is None:
                nivel.append(None); continue
            if prev[j] is not None:
                d = t[j] - prev[j]
                if d < 0: cons[j] = round(-d * fac, 3)
                elif d > 0:
                    rec += d * fac
                    if j == 2: sube3 = d * fac
            prev[j] = t[j]
            nivel.append(round(t[j] * fac, 2))
        # Traspaso TK4 -> TK3 (habitual desde el cambio de fórmula del TK3): si el mismo día baja el TK4
        # y sube el TK3, lo que subió el TK3 es ácido traspasado y no se cuenta como consumo ni como recepción.
        tras = 0.0
        if f > CAMBIO_TK3 and cons[3] > 0 and sube3 > 0:
            tras = min(cons[3], sube3)
            cons[3] = round(cons[3] - tras, 3); rec -= tras
        dias.append([f, round(sum(cons), 3), [round(c, 3) for c in cons], round(max(rec, 0), 2), nivel, round(tras, 2)])
    # Días con llegada de camión: la regla no cuenta consumo ese día (la lectura sube), pero la planta sí consumió.
    # Se estima el consumo de esos días con el promedio de los 3 días anteriores y 3 posteriores sin llegada.
    for k, d in enumerate(dias):
        d.append(0.0)
        if d[3] <= 0: continue
        prev = [x[1] for x in dias[max(0, k-15):k] if x[3] <= 0][-3:]
        nxt = [x[1] for x in dias[k+1:k+16] if x[3] <= 0][:3]
        vals = prev + nxt
        est = sum(vals) / len(vals) if vals else 0.0
        if est > d[1]:
            d[6] = round(est - d[1], 3)
            d[1] = round(est, 3)
    return dias, avisos

def main(contenido=None):
    dias, avisos = leer(contenido or descargar(URL))
    if len(dias) < 20: sys.exit(f"Muy pocos días ({len(dias)}): no se actualiza.")
    datos = {"actualizado": max(d[0] for d in dias), "cambioTK3": CAMBIO_TK3, "dias": dias, "avisos": avisos}
    nuevo = json.dumps(datos, separators=(",", ":"), ensure_ascii=False)
    viejo = open(SALIDA, encoding="utf-8").read() if os.path.exists(SALIDA) else ""
    if nuevo == viejo: print("Sin cambios."); return
    open(SALIDA, "w", encoding="utf-8").write(nuevo)
    print(f"acido.json: {len(dias)} lecturas; fechas corregidas {avisos['fechas']}; lecturas con coma corregida {avisos['comas']}")

if __name__ == "__main__":
    main(open(sys.argv[1], "rb").read() if len(sys.argv) > 1 else None)
