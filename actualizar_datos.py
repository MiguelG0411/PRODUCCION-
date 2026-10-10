"""Descarga el Excel de producción desde Dropbox y genera datos.json para la app.
Lo ejecuta GitHub Actions de forma automática (ver .github/workflows/actualizar.yml)."""
import io, json, os, sys, urllib.request, datetime
from openpyxl import load_workbook

URL = os.environ.get("DROPBOX_URL", "").strip()
SALIDA = "datos.json"
KG_SACO = 25

def descargar(url):
    if not url:
        sys.exit("Falta el secreto DROPBOX_URL en GitHub.")
    # Forzar descarga directa del archivo en los enlaces de Dropbox
    url = url.replace("dl=0", "dl=1")
    if "dropbox.com" in url and "dl=1" not in url:
        url += ("&" if "?" in url else "?") + "dl=1"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()

def num(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return 0

def leer(contenido):
    wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    if "BDR" not in wb.sheetnames:
        sys.exit("No se encontró la hoja BDR en el Excel.")
    dias = {}
    # Hoja BDR: desde la fila 6, columnas C (día), E (sacos turno día), F (sacos turno noche)
    for fila in wb["BDR"].iter_rows(min_row=6, min_col=3, max_col=6, values_only=True):
        fecha, _mes, sd, sn = fila
        if isinstance(fecha, datetime.datetime):
            fecha = fecha.date()
        if not isinstance(fecha, datetime.date):
            continue
        k = fecha.isoformat()
        a, b = dias.get(k, [0, 0])  # si un día viene repetido, se suman
        dias[k] = [a + num(sd), b + num(sn)]
    return dias

def main():
    dias = leer(descargar(URL))
    if len(dias) < 10:
        sys.exit(f"Muy pocos días leídos ({len(dias)}): no se actualiza para no borrar datos.")
    datos = {"kgSaco": KG_SACO, "actualizado": max(dias), "dias": dict(sorted(dias.items()))}
    nuevo = json.dumps(datos, separators=(",", ":"), ensure_ascii=False)
    viejo = open(SALIDA, encoding="utf-8").read() if os.path.exists(SALIDA) else ""
    if nuevo == viejo:
        print("Sin cambios en la producción.")
        return
    open(SALIDA, "w", encoding="utf-8").write(nuevo)
    total = sum((a + b) * KG_SACO for a, b in dias.values())
    print(f"datos.json actualizado: {len(dias)} días, {total:,} kg, último día {max(dias)}")

if __name__ == "__main__":
    main()
