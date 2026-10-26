"""Funciones compartidas por los lectores de planillas."""
import datetime, io, urllib.request
from openpyxl import load_workbook


class ErrorPlanilla(Exception):
    """La planilla no tiene el formato esperado; se conservan los datos anteriores."""


def descargar(url, timeout=120):
    # Los links compartidos de Dropbox con dl=1 descargan el archivo directo
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def libro(contenido):
    # data_only=True: valores calculados, no fórmulas
    return load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)


def hoja(wb, nombre):
    if nombre not in wb.sheetnames:
        raise ErrorPlanilla(f"No se encontró la hoja {nombre}.")
    return wb[nombre]


def filas(ws, desde_col, hasta_col, desde_fila=1):
    """Recorre las filas entregando (n° de fila, valores) con largo fijo aunque la fila venga corta."""
    n = hasta_col - desde_col + 1
    for i, r in enumerate(ws.iter_rows(min_row=desde_fila, min_col=desde_col, max_col=hasta_col, values_only=True), start=desde_fila):
        r = tuple(r)
        yield i, r + (None,) * (n - len(r))


def fecha(v):
    if isinstance(v, datetime.datetime): v = v.date()
    if isinstance(v, datetime.date) and v.year >= 2000: return v.isoformat()
    return None


def num(v):
    if isinstance(v, bool): return None
    try: return float(v)
    except (TypeError, ValueError): return None


def txt(v):
    return " ".join(str(v or "").upper().split())


def minimo(n, minimo_, que):
    # Evita reemplazar los datos buenos por una lectura casi vacía (archivo dañado o formato cambiado)
    if n < minimo_:
        raise ErrorPlanilla(f"Muy pocos {que} ({n}); se mantienen los datos anteriores.")
